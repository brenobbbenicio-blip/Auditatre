"""Leitura local de PDF, planilha, CSV, DOCX e HTML. Sem rede."""

from __future__ import annotations

import csv
import io
import zipfile
from xml.etree import ElementTree

from auditatre.extract import html_para_texto, planilha_para_texto


def ler_texto(caminho: str, dados: bytes) -> tuple[str, str]:
    nome = caminho.lower()
    if nome.endswith(".pdf"):
        return _pdf(dados)
    if nome.endswith(".docx"):
        return _docx(dados), "extraido"
    if nome.endswith(".xlsx"):
        return _xlsx(dados), "extraido"
    if nome.endswith(".csv"):
        return _csv(dados), "extraido"
    if nome.endswith(".html") or nome.endswith(".htm"):
        return html_para_texto(dados.decode("utf-8", "replace")), "extraido"
    if nome.endswith(".txt") or nome.endswith(".md"):
        return dados.decode("utf-8", "replace"), "extraido"
    return "", "nao_lido"


def _pdf(dados: bytes) -> tuple[str, str]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(dados))
        if reader.is_encrypted:
            return "", "sem_texto"
        partes = [(pagina.extract_text() or "") for pagina in reader.pages]
    except (PdfReadError, ValueError, OSError):
        return "", "sem_texto"
    texto = "\n".join(partes).strip()
    if not texto:
        return "", "sem_texto"
    return texto, "extraido"


def _docx(dados: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(dados)) as arquivo:
        xml = arquivo.read("word/document.xml")
    raiz = ElementTree.fromstring(xml)
    return " ".join(trecho.strip() for trecho in raiz.itertext() if trecho.strip())


def _xlsx(dados: bytes) -> str:
    from openpyxl import load_workbook

    livro = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    blocos: list[str] = []
    for aba in livro.worksheets:
        linhas: list[list[str]] = []
        for linha in aba.iter_rows(values_only=True):
            linhas.append(["" if celula is None else str(celula) for celula in linha])
        texto = planilha_para_texto(linhas)
        if texto:
            blocos.append(f"# {aba.title}\n{texto}")
    livro.close()
    return "\n".join(blocos)


def _csv(dados: bytes) -> str:
    texto = dados.decode("utf-8-sig", "replace")
    leitor = csv.reader(io.StringIO(texto))
    return planilha_para_texto([linha for linha in leitor])
